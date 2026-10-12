"""Эталонные значения: пересчёт КОПИИ книги самим Excel (COM) при заполненных RPE.

Исходный файл не открывается на запись: копия кладётся во временную папку.
Результат — tests/fixtures/excel-golden.json: входы каждого сценария и все
рассчитанные Excel значения плана (подготовительный период и 12 недель).

Нужен Windows + Microsoft Excel + pywin32.
Запуск: <python с pywin32> scripts/excel_golden.py
"""
import json
import shutil
import tempfile
from pathlib import Path

import openpyxl
import win32com.client
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "Жимовая прога v 2.0.xlsx"
OUT = ROOT / "tests" / "fixtures" / "excel-golden.json"

# точные названия из столбца L листа «Формулы» (с хвостовыми пробелами — VLOOKUP точный)
names = {}
wb = openpyxl.load_workbook(SRC, read_only=True)
for row in wb["Формулы"].iter_rows(min_row=9, max_row=69, min_col=12, max_col=12):
    v = row[0].value
    if isinstance(v, str) and v.strip() and v.strip() != "+":
        names.setdefault(v.strip(), v)
wb.close()
N = lambda s: names[s]

SLOT_CELLS = {  # слот -> ячейка выбора на подготовительном периоде, строка теста, строка мезоцикла 1
    "d1s1": ("B77", 111, 154), "d1s2": ("B78", 112, 155), "d1s3": ("B79", 113, 156), "d1s4": ("B80", 114, 157), "d1s5": ("B81", 115, 158),
    "d2s1": ("B85", 119, 162), "d2s2": ("B86", 120, 163), "d2s3": ("B87", 121, 164), "d2s4": ("B88", 122, 165), "d2s5": ("B89", 123, 166),
    "d3s1": ("B93", 127, 170), "d3s2": ("B94", 128, 171), "d3s3": ("B95", 129, 172), "d3s4": ("B96", 130, 173), "d3s5": ("B97", 131, 174),
}
RPE_CYCLE = [8, 8.5, 9, 7.5, 9.5, 7, 6.5, 10, 8.5, 9, 7.5, 8]

SCENARIOS = {
    "A_исходный_выбор": {"choices": {}, "tests": {115: (None, None), 123: (None, None), 131: (None, None)}},
    "B_альтернативы": {
        "choices": {"d1s1": N("Жим гантелей лёжа"), "d1s2": N("Подтягивания (собственный вес)"), "d1s3": N("Жим платформы ногами"),
                    "d1s4": N("Тяга к подбородку"), "d1s5": N("Скручивания"), "d2s2": N("Мёртвая тяга"), "d2s3": N("Молот"),
                    "d2s4": N("Разгибания рук на блоке"), "d2s5": N("Шраги"), "d3s1": N("Приседания в колодец"),
                    "d3s2": N("Армейский жим"), "d3s3": N("Тяга гантели в наклоне"), "d3s4": N("Пуловер"), "d3s5": N("Сгибания ног")},
        "tests": {111: (30, 8), 112: (None, 12), 113: (180, 8), 114: (40, 8), 115: (None, 20), 120: (120, 6), 121: (20, 8),
                  122: (35, 10), 123: (100, 8), 127: (100, 5), 128: (55, 5), 129: (32, 8), 130: (30, 10), 131: (40, 10)},
    },
    "C_альтернативы_2": {
        "choices": {"d1s1": N("Жим штанги лёжа на наклонной скамье"), "d1s2": N("Тяга вертикального блока"), "d1s3": N("Фронтальные приседания"),
                    "d1s5": N("Экстензия"), "d2s2": N("Ягодичный мост"), "d2s4": N("Жим лёжа узким хватом"), "d2s5": N("Подъёмы на носки"),
                    "d3s1": N("Ягодичный мост"), "d3s3": N("Тяга штанги в наклоне"), "d3s4": N("Тяга горизонтального блока к поясу"),
                    "d3s5": N("Разгибания ног")},
        "tests": {111: (70, 5), 112: (75, 8), 113: (90, 5), 115: (None, 15), 120: (120, 8), 122: (85, 5), 123: (80, 12),
                  127: (110, 6), 129: (70, 6), 130: (60, 10), 131: (45, 10)},
    },
}


def rpe_cells():
    """Ячейки ввода RPE (F, M, T каждого мезоцикла) и спецячейки жима."""
    cells = []
    for _, (_, _, r1) in SLOT_CELLS.items():
        for meso in range(3):
            r = r1 + 27 * meso
            cells += [f"F{r}", f"M{r}", f"T{r}"]
    cells += ["X162", "X189", "X216", "AI208", "AI216"]
    return cells


excel = win32com.client.DispatchEx("Excel.Application")
excel.Visible = False
excel.DisplayAlerts = False
result = {"source": SRC.name, "rpeCycle": RPE_CYCLE, "scenarios": {}}
tmp = Path(tempfile.mkdtemp())
try:
    for sname, sc in SCENARIOS.items():
        copy = tmp / f"golden_{len(result['scenarios'])}.xlsx"
        shutil.copy2(SRC, copy)
        book = excel.Workbooks.Open(str(copy))
        ws = book.Worksheets("Программа тренировок ")
        for slot, nm in sc["choices"].items():
            ws.Range(SLOT_CELLS[slot][0]).Value = nm
        for r, (w, reps) in sc["tests"].items():
            ws.Range(f"C{r}").Value = w
            ws.Range(f"D{r}").Value = reps
        ws.Range("K80").Value = 8.5  # пустое RPE подготовительного периода в исходнике
        for i, c in enumerate(rpe_cells()):
            ws.Range(c).Value = RPE_CYCLE[i % len(RPE_CYCLE)]
        ws.Range("AH162").Value = 4  # повторения AMRAP недели 4 вместо «?»
        ws.Range("AH189").Value = 3  # повторения AMRAP недели 8 вместо «?»
        excel.CalculateFull()
        cells = {}
        for r in list(range(77, 98)) + list(range(111, 132)) + list(range(154, 229)):
            for c in range(2, ci("AP") + 1):
                v = ws.Cells(r, c).Value
                if v is None or v == "":
                    continue
                if isinstance(v, int) and v < -2_000_000_000:  # код ошибки COM
                    v = "#ERR"
                cells[f"{gl(c)}{r}"] = v
        result["scenarios"][sname] = {"choices": sc["choices"], "cells": cells}
        book.Close(SaveChanges=False)
finally:
    excel.Quit()
    shutil.rmtree(tmp, ignore_errors=True)

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
print("scenarios:", list(result["scenarios"]), "->", OUT)
